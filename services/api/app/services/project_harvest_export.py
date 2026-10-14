"""项目收获分析 Excel 异步导出（Redis 存任务状态，文件放 OSS，24h 过期）。数据量小（地块×日期），线程池生成。"""

from __future__ import annotations

import asyncio
import io
import json
import time
import uuid
from collections.abc import Mapping, Sequence
from datetime import datetime, timezone
from typing import Any

import structlog
from agric_satellite_analysis_common import project_harvest as ph
from starlette.concurrency import run_in_threadpool

from app.services import project_harvest_views as v

logger = structlog.get_logger(__name__)

JOB_TTL_SEC = 24 * 3600
STALE_SEC = 300
_KEY = "project-harvest-export:job:{}"
_tasks: set[asyncio.Task] = set()
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
STATUS_LABEL = {"1": "在用(1)", "0": "停用(0)"}


def build_workbook(
    group_id: str,
    run: Mapping[str, Any],
    stats: Sequence[Mapping[str, Any]],
    summaries: Sequence[Mapping[str, Any]],
    dedup: Sequence[Mapping[str, Any]],
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    meta = {str(d["land_id"]): d for d in dedup}
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    names = [t["name"] for t in ph.TIERS]
    ws = wb.active
    ws.title = "地块明细"
    rows = []
    for r in v.parcel_rows(stats, meta, "pick"):
        rows.append([r["land_id"], r["land_name"], STATUS_LABEL.get(str(r["land_status"]), r["land_status"]), r["town_name"],
                     r["village_name"], r["block_key"], r["area_mu"], r["crop_mu"]]
                    + [r[f"{k}_mu"] for k in ph.TIER_KEYS]
                    + [r["harvested_pct"], r["cloud_pct"], r["peak_ndvi_median"], r["obs_date"], "回退" if r["fallback"] else "",
                       r["registered_area_mu"], ",".join(r["covered_by"])])
    sheet(ws, ["land_id", "地块名称", "状态", "乡镇", "村", "计算块", "去重后实测亩", "耕种亩"] + [f"{n}亩" for n in names]
          + ["已收获(高+中)%", "云遮%", "峰值NDVI中位数", "采用日期", "说明", "原登记亩", "被谁覆盖"], rows)
    ws2 = wb.create_sheet("汇总")
    pick = v.summary_view(summaries, "pick", str(run.get("season_to")), (run.get("extras") or {}).get("pick_dates"))
    rows = []
    for title, items in (("全部", [{"key": "全部", **(pick["totals"] or {})}]), ("按计算块", pick["by_block"]),
                         ("按乡镇/村", pick["by_town_village"]), ("按状态", pick["by_status"])):
        for it in items:
            rows.append([title, it["key"], it.get("parcel_count"), it.get("area_mu"), it.get("crop_mu")]
                        + [it.get(f"{k}_mu") for k in ph.TIER_KEYS] + [it.get("harvested_pct")])
    sheet(ws2, ["维度", "分组", "地块数", "去重实测亩", "耕种亩"] + [f"{n}亩" for n in names] + ["已收获(高+中)%"], rows)
    ws3 = wb.create_sheet("日期序列")
    sheet(ws3, ["日期", "已收获%(占耕种)", "云遮%", "耕种亩(有影像)"],
          [[s["date"], s.get("harvested_pct"), s.get("cloud_pct"), s.get("crop_mu")] for s in pick["series"]])
    ws4 = wb.create_sheet("被合并地块")
    sheet(ws4, ["land_id", "地块名称", "状态", "原因", "原边界实测亩", "剩余亩", "被谁覆盖", "优先级序号"],
          [[d["land_id"], d.get("land_name"), STATUS_LABEL.get(str(d.get("land_status")), d.get("land_status")), d.get("merged_reason"),
            d.get("original_area_mu"), d.get("clipped_area_mu"), ",".join(d.get("covered_by") or []), d.get("priority_rank")]
           for d in v.dedup_view(dedup)["items"] if not d["kept"]])
    ws5 = wb.create_sheet("说明")
    p = run.get("params") or {}
    for line in [
        f"项目 {group_id}  算法 {ph.METHOD_VERSION}  运行 {run.get('run_id')}  季节 {run.get('season_from')}~{run.get('season_to')}",
        f"耕种判定：≤{p.get('early_cutoff')} 晴空峰值NDVI ≥ {p.get('green_thr')}；收获：NDVI<{p.get('low_thr')} 且较峰值降≥{p.get('drop_min')}，其后晴空观测均<{p.get('stay_thr')}。",
        "去重：优先级=(状态 启用1>停用0, land_id 大者优先)；每块=原边界−更高优先级边界并集；剩余<0.3亩或比1像元细的记为被合并。",
        "面积：去重后实测亩；各分级亩=覆盖比例加权像元比例×实测亩；%以耕种面积为分母。采用日期=耕种部分云遮≤20%的最新日期。",
    ]:
        ws5.append([line])
    ws5.column_dimensions["A"].width = 140
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def get_job(redis, job_id: str) -> dict | None:
    raw = await redis.get(_KEY.format(job_id))
    if not raw:
        return None
    job = json.loads(raw)
    if job["status"] in ("pending", "running") and time.time() - job.get("heartbeat_at", 0) > STALE_SEC:
        job.update(status="failed", error="导出任务所在实例已中断，请重新导出")
    return job


async def _save(redis, job: dict) -> None:
    job["heartbeat_at"] = time.time()
    await redis.set(_KEY.format(job["job_id"]), json.dumps(job), ex=JOB_TTL_SEC)


def public_view(job: Mapping[str, Any]) -> dict[str, Any]:
    return {k: job.get(k) for k in ("job_id", "group_id", "status", "error", "file_name", "size_bytes", "created_at", "finished_at")}


async def create_job(redis, group_id: str, loader) -> dict:
    """loader: async () -> (run, stats, summaries, dedup)。"""
    job = {"job_id": uuid.uuid4().hex, "group_id": group_id, "status": "pending", "error": None,
           "file_name": f"project-harvest_{group_id}_{datetime.now(timezone.utc):%Y%m%d}.xlsx",
           "object_key": None, "size_bytes": None, "created_at": time.time(), "finished_at": None}
    await _save(redis, job)

    async def runner():
        try:
            await _save(redis, {**job, "status": "running"})
            job["status"] = "running"
            run, stats, summaries, dedup = await loader()
            data = await run_in_threadpool(build_workbook, group_id, run, stats, summaries, dedup)
            from app.core.storage import get_storage

            key = f"exports/project-harvest/{datetime.now(timezone.utc):%Y%m%d}/{job['job_id']}.xlsx"
            await run_in_threadpool(get_storage().put_bytes, key, data, content_type=XLSX_MIME)
            job.update(status="done", object_key=key, size_bytes=len(data), finished_at=time.time())
        except Exception as exc:
            logger.exception("project_harvest_export_failed", job_id=job["job_id"])
            job.update(status="failed", error=str(exc)[:500], finished_at=time.time())
        await _save(redis, job)

    t = asyncio.create_task(runner())
    _tasks.add(t)
    t.add_done_callback(_tasks.discard)
    return job
