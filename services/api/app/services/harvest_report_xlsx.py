"""收获进度报表 Excel（openpyxl）：汇总 / 逐日明细 / 透视 / 说明 四个工作表。"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.harvest_report import (
    BUCKET_LABELS_ZH,
    BUCKETS,
    CONFIDENCE_LABELS_ZH,
    STATUS_LABELS_ZH,
    ReportFilters,
)

CN_TZ = timezone(timedelta(hours=8))

HEADER_FILL = PatternFill("solid", fgColor="1F6F43")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
KPI_LABEL_FILL = PatternFill("solid", fgColor="E8F3EC")
CARRY_FONT = Font(italic=True)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BUCKET_FILLS = {
    "none": "F3F4F6",
    "lt30": "FEF3C7",
    "30_90": "FDBA74",
    "ge90": "C2410C",
    "not_computed": "E5E7EB",
}

PCT = '0.0"%"'
MU = "#,##0.00"
DATE = "yyyy-mm-dd"

# (key, 表头, 列宽, 数字格式)
SUMMARY_COLS: list[tuple[str, str, int, str | None]] = [
    ("land_id", "地块ID", 10, None),
    ("land_name", "地块名称", 24, None),
    ("group_name", "分组/农场", 18, None),
    ("crop_type", "作物", 8, None),
    ("province_name", "省", 12, None),
    ("city_name", "市", 10, None),
    ("county_name", "区县", 10, None),
    ("area_mu", "面积(亩)", 11, MU),
    ("season_start", "本季起始日", 12, DATE),
    ("first_harvest_date", "首次收获日", 12, DATE),
    ("latest_date", "最新结果日", 12, DATE),
    ("latest_obs_date", "最新影像日", 12, DATE),
    ("days_since_last_image", "距最新影像(天)", 10, "0"),
    ("harvested_pct", "已收获%", 10, PCT),
    ("suspected_pct", "疑似收获%", 10, PCT),
    ("combined_pct", "合计%(已收获+疑似)", 14, PCT),
    ("harvested_area_mu", "已收获面积(亩)", 12, MU),
    ("combined_area_mu", "合计收获面积(亩)", 13, MU),
    ("newly_combined_pct", "较上期新增%", 11, PCT),
    ("status_zh", "状态", 9, None),
    ("bucket_zh", "进度分档", 14, None),
    ("confidence", "置信度", 8, "0.00"),
    ("confidence_zh", "置信等级", 8, None),
    ("confirmed_zh", "已确认", 7, None),
    ("obs_count", "区间结果期数", 9, "0"),
    ("method_version", "算法版本", 26, None),
]

DAILY_COLS: list[tuple[str, str, int, str | None]] = [
    ("land_id", "地块ID", 10, None),
    ("land_name", "地块名称", 24, None),
    ("group_name", "分组/农场", 18, None),
    ("crop_type", "作物", 8, None),
    ("county_name", "区县", 10, None),
    ("area_mu", "面积(亩)", 11, MU),
    ("date", "观测日", 12, DATE),
    ("status_zh", "状态", 9, None),
    ("harvested_pct", "已收获%", 10, PCT),
    ("suspected_pct", "疑似收获%", 10, PCT),
    ("combined_pct", "合计%", 10, PCT),
    ("newly_harvested_pct", "新增已收获%", 11, PCT),
    ("newly_combined_pct", "较上期新增合计%", 13, PCT),
    ("harvested_area_mu", "已收获面积(亩)", 12, MU),
    ("combined_area_mu", "合计收获面积(亩)", 13, MU),
    ("valid_pct", "有效像元%", 10, PCT),
    ("season_start", "本季起始日", 12, DATE),
    ("confidence", "置信度", 8, "0.00"),
    ("confidence_zh", "置信等级", 8, None),
    ("confirmed_zh", "已确认", 7, None),
    ("method_version", "算法版本", 26, None),
]


def _decorate(r: dict[str, Any]) -> dict[str, Any]:
    out = dict(r)
    out["status_zh"] = STATUS_LABELS_ZH.get(
        r.get("status") or "", r.get("status") or ""
    )
    if "bucket" in r:
        out["bucket_zh"] = BUCKET_LABELS_ZH.get(r["bucket"], r["bucket"])
    out["confidence_zh"] = CONFIDENCE_LABELS_ZH.get(r.get("confidence_level") or "", "")
    c = r.get("confirmed")
    out["confirmed_zh"] = "" if c is None else ("是" if c else "否（暂定）")
    return out


def _header(ws, row: int, cols) -> None:
    for i, (_k, title, width, _fmt) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 32


def _rows(ws, start: int, cols, records) -> int:
    r = start
    for rec in records:
        for i, (k, _t, _w, fmt) in enumerate(cols, start=1):
            v = rec.get(k)
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if fmt and v is not None:
                c.number_format = fmt
        r += 1
    return r - 1


def _print_setup(ws, title_rows: str) -> None:
    """打印：横向、按页宽缩放、每页重复表头。"""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows


def _col(cols, key: str) -> str:
    return get_column_letter([c[0] for c in cols].index(key) + 1)


def _pct_bars(ws, cols, keys, first: int, last: int) -> None:
    if last < first:
        return
    for key, color in keys:
        rng = f"{_col(cols, key)}{first}:{_col(cols, key)}{last}"
        ws.conditional_formatting.add(
            rng,
            DataBarRule(
                start_type="num",
                start_value=0,
                end_type="num",
                end_value=100,
                color=color,
            ),
        )


def _filters_text(f: ReportFilters) -> list[tuple[str, str]]:
    return [
        ("日期区间", f"{f.date_from.isoformat()} ~ {f.date_to.isoformat()}"),
        ("分组ID", f.group_id or "全部"),
        ("作物", f.crop or "全部"),
        (
            "进度分档",
            "、".join(BUCKET_LABELS_ZH.get(b, b) for b in f.buckets) or "全部",
        ),
        ("最低合计%", "" if f.min_pct is None else f"{f.min_pct:g}"),
        ("关键字", f.keyword or ""),
    ]


def build_workbook(
    parcels: list[dict[str, Any]],
    daily: list[dict[str, Any]],
    kpi: dict[str, Any],
    filters: ReportFilters,
    generated_at: datetime | None = None,
    carry_forward: bool = True,
) -> bytes:
    gen = (generated_at or datetime.now(CN_TZ)).astimezone(CN_TZ)
    wb = Workbook()

    # ---- Sheet1 汇总 ----
    ws = wb.active
    ws.title = "汇总"
    ws["A1"] = "地块收获进度汇总"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"区间 {filters.date_from.isoformat()} ~ {filters.date_to.isoformat()}"
        f"　生成时间 {gen:%Y-%m-%d %H:%M} (UTC+8)"
    )
    ws["A2"].font = Font(color="6B7280")
    bc = kpi["bucket_counts"]
    kpis = [
        ("地块数", kpi["parcel_count"], "0"),
        ("已计算地块数", kpi["computed_count"], "0"),
        ("总面积(亩)", kpi["total_area_mu"], MU),
        ("已收获面积(亩)", kpi["harvested_area_mu"], MU),
        ("合计收获面积(亩)", kpi["combined_area_mu"], MU),
        ("平均合计%", kpi["avg_combined_pct"], PCT),
        ("面积加权合计%", kpi["area_weighted_combined_pct"], PCT),
    ] + [(BUCKET_LABELS_ZH[b], bc.get(b, 0), "0") for b in BUCKETS]
    # 两行 KPI：标签行 + 数值行
    for i, (label, value, fmt) in enumerate(kpis, start=1):
        lc = ws.cell(row=4, column=i, value=label)
        lc.fill, lc.font, lc.border = KPI_LABEL_FILL, Font(bold=True, size=9), BORDER
        lc.alignment = Alignment(horizontal="center", wrap_text=True)
        vc = ws.cell(row=5, column=i, value=value)
        vc.font, vc.border, vc.number_format = Font(bold=True, size=12), BORDER, fmt
        vc.alignment = Alignment(horizontal="center")
    ws.row_dimensions[4].height = 30
    hdr = 7
    _header(ws, hdr, SUMMARY_COLS)
    last = _rows(ws, hdr + 1, SUMMARY_COLS, (_decorate(p) for p in parcels))
    ws.freeze_panes = ws.cell(row=hdr + 1, column=3)
    ws.auto_filter.ref = (
        f"A{hdr}:{get_column_letter(len(SUMMARY_COLS))}{max(last, hdr)}"
    )
    _pct_bars(
        ws,
        SUMMARY_COLS,
        [
            ("harvested_pct", "C2410C"),
            ("suspected_pct", "FDBA74"),
            ("combined_pct", "EA580C"),
        ],
        hdr + 1,
        last,
    )
    bcol = _col(SUMMARY_COLS, "bucket_zh")
    for r, p in enumerate(parcels, start=hdr + 1):
        fill = BUCKET_FILLS.get(p["bucket"])
        if fill:
            cell = ws[f"{bcol}{r}"]
            cell.fill = PatternFill("solid", fgColor=fill)
            if p["bucket"] == "ge90":
                cell.font = Font(color="FFFFFF", bold=True)

    _print_setup(ws, f"{hdr}:{hdr}")

    # ---- Sheet2 逐日明细 ----
    ws2 = wb.create_sheet("逐日明细")
    _header(ws2, 1, DAILY_COLS)
    recs = sorted(daily, key=lambda r: (r["land_id"], r["date"] or date.min))
    last2 = _rows(ws2, 2, DAILY_COLS, (_decorate(r) for r in recs))
    ws2.freeze_panes = "C2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(DAILY_COLS))}{max(last2, 1)}"
    _pct_bars(ws2, DAILY_COLS, [("combined_pct", "EA580C")], 2, last2)

    _print_setup(ws2, "1:1")

    # ---- Sheet3 透视 ----
    ws3 = wb.create_sheet("透视")
    dates = sorted({r["date"] for r in daily if r["date"]})
    fixed = [
        ("land_id", "地块ID", 10),
        ("land_name", "地块名称", 24),
        ("area_mu", "面积(亩)", 10),
    ]
    for i, (_k, title, width) in enumerate(fixed, start=1):
        c = ws3.cell(row=1, column=i, value=title)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        ws3.column_dimensions[get_column_letter(i)].width = width
    for j, d in enumerate(dates, start=len(fixed) + 1):
        c = ws3.cell(row=1, column=j, value=d)
        c.number_format = "mm-dd"
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", text_rotation=90)
        ws3.column_dimensions[get_column_letter(j)].width = 6
    ws3.row_dimensions[1].height = 48
    cell_of: dict[tuple[str, date], dict[str, Any]] = {
        (r["land_id"], r["date"]): r for r in daily if r["date"]
    }
    for i, p in enumerate(parcels, start=2):
        ws3.cell(row=i, column=1, value=p["land_id"])
        ws3.cell(row=i, column=2, value=p["land_name"])
        a = ws3.cell(row=i, column=3, value=p["area_mu"])
        a.number_format = MU
        carried: dict[str, Any] | None = None
        for j, d in enumerate(dates, start=len(fixed) + 1):
            r = cell_of.get((p["land_id"], d))
            if r is not None:
                carried = r
                c = ws3.cell(row=i, column=j, value=r["combined_pct"])
                c.number_format = "0"
            elif (
                carry_forward
                and carried is not None
                and (
                    p["season_start"] is None
                    or carried["season_start"] == p["season_start"]
                )
            ):
                c = ws3.cell(row=i, column=j, value=carried["combined_pct"])
                c.number_format = "0"
                c.font = CARRY_FONT
    if dates and parcels:
        rng = (
            f"{get_column_letter(len(fixed) + 1)}2:"
            f"{get_column_letter(len(fixed) + len(dates))}{len(parcels) + 1}"
        )
        ws3.conditional_formatting.add(
            rng,
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="F0FDF4",
                mid_type="num",
                mid_value=50,
                mid_color="FDBA74",
                end_type="num",
                end_value=100,
                end_color="9A3412",
            ),
        )
        # 深色单元格用白字，保证沿用值（斜体）与实测值都可读
        ws3.conditional_formatting.add(
            rng,
            CellIsRule(
                operator="greaterThanOrEqual", formula=["60"], font=Font(color="FFFFFF")
            ),
        )
    ws3.freeze_panes = "D2"
    ws3.auto_filter.ref = (
        f"A1:{get_column_letter(len(fixed) + len(dates))}{len(parcels) + 1}"
    )

    _print_setup(ws3, "1:1")

    # ---- Sheet4 说明 ----
    ws4 = wb.create_sheet("说明")
    ws4.column_dimensions["A"].width = 22
    ws4.column_dimensions["B"].width = 100
    lines: list[tuple[str, str]] = [
        ("生成时间", f"{gen:%Y-%m-%d %H:%M:%S} (UTC+8)"),
        *_filters_text(filters),
        ("", ""),
        (
            "算法",
            "Sentinel-2 残茬/裸土单调判别 + Sentinel-1 佐证（"
            f"{parcels[0]['method_version'] if parcels and parcels[0]['method_version'] else 's2s1_residue_monotonic_v4'}"
            "）：每像元在本季内状态单调（未收获→疑似收获→已收获），占比只增不减。",
        ),
        (
            "落库门槛",
            f"只有合计占比 > {kpi.get('min_save_pct', 3):g}% 的观测日入库；"
            "区间内没有结果行但已计算区间标记覆盖查询区间的地块，记为“未收获（≤门槛）”，合计 0%。",
        ),
        (
            "未计算",
            "既无结果行也无已计算区间标记的地块（尚未运行 harvest-progress 计算/回填），"
            "不计入平均值与收获面积；报表不做现场计算。",
        ),
        (
            "透视表",
            "单元格为该观测日合计%（已收获+疑似）。"
            + (
                "斜体为沿用上一期结果（结果单调，未入库日沿用上一期值，仅限同一季）；"
                "首个结果日之前留空（≤门槛）。"
                if carry_forward
                else "仅显示入库观测日，未入库日留空。"
            ),
        ),
        ("", ""),
        ("字段", "定义"),
        ("已收获%", "高确定性已收获像元占作物像元比例（单调）。"),
        (
            "疑似收获%",
            "疑似收获像元占比（出现残茬/突变特征但尚未确认）；后续确认即转为已收获。",
        ),
        ("合计%", "已收获% + 疑似收获%（单调）；进度分档与“最低合计%”筛选均按此值。"),
        ("已收获面积(亩)", "面积 × 已收获%。"),
        ("合计收获面积(亩)", "面积 × 合计%。"),
        (
            "较上期新增%",
            "合计% 相对同一季上一条结果行的增量；无上一条时即本期合计（此前均 ≤ 门槛）。",
        ),
        ("新增已收获%", "算法输出的已收获% 较上一观测日新增。"),
        ("首次收获日", "本季第一个合计% 超过落库门槛的观测日。"),
        ("最新结果日", "查询区间内最后一个入库观测日；地块最新状态取自该日。"),
        (
            "最新影像日 / 距最新影像(天)",
            "截至区间末日最近一期 S2 影像日期，以及到 min(区间末日, 今天) 的天数。",
        ),
        ("状态", "生长中 / 收获中 / 已收获（算法状态）。"),
        ("进度分档", "未收获（≤门槛）/ 收获 <30% / 收获 30–90% / 收获 ≥90% / 未计算。"),
        (
            "置信度 / 置信等级",
            "0–1 分值与 高/中/低 等级（像元数、观测间隔、S1 一致性等）。",
        ),
        ("已确认", "否（暂定）表示最近观测尚未被后续观测确认，后续可能修正。"),
        ("区间结果期数", "查询区间内入库观测日数。"),
    ]
    for i, (k, v) in enumerate(lines, start=1):
        a, b = ws4.cell(row=i, column=1, value=k), ws4.cell(row=i, column=2, value=v)
        a.font = Font(bold=True)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if k == "字段":
            a.fill = b.fill = HEADER_FILL
            a.font = b.font = HEADER_FONT

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
