"""收获报表异步导出任务：创建/复用/心跳超时/失败/下载/清理。"""

from __future__ import annotations

import asyncio
import io
import time
import unittest
from datetime import date, datetime, timedelta, timezone
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services import harvest_report as hr
from app.services import harvest_report_export as hx

F = hr.ReportFilters(date_from=date(2026, 8, 1), date_to=date(2026, 10, 10))


class FakeRedis:
    def __init__(self):
        self.d: dict[str, str] = {}

    async def get(self, k):
        return self.d.get(k)

    async def set(self, k, v, ex=None):
        self.d[k] = v


class FakeStorage:
    def __init__(self):
        self.objs: dict[str, bytes] = {}
        self.deleted: list[str] = []

    def upload_file(self, key, path, content_type=None):
        with open(path, "rb") as fh:
            self.objs[key] = fh.read()
        return key

    def iter_bytes(self, key, chunk_size=65536):
        data = self.objs[key]
        for i in range(0, len(data), chunk_size):
            yield data[i : i + chunk_size]

    def list_keys(self, prefix, suffix="", limit=0):
        return [k for k in self.objs if k.startswith(prefix) and k.endswith(suffix)]

    def delete(self, key):
        self.deleted.append(key)
        self.objs.pop(key, None)


def _land(lid):
    return {"land_id": lid, "land_name": f"地块{lid}", "group_id": "1", "group_name": "组",
            "crop_type": "corn", "area_mu": 100.0, "province_name": "吉林省", "base_name": "租户"}



async def _inline_build(rows, daily, kpi, filters, cf, path):
    return hx.build_to_file(rows, daily, kpi, filters, cf, path)


def _params(**kw):
    return hx.params_of(F, kw.get("sort", "combined_pct"), "desc", True)


class ExportJobTests(unittest.TestCase):
    def setUp(self):
        self.redis = FakeRedis()
        self.storage = FakeStorage()
        self.patches = [
            patch.object(hx, "_load", side_effect=_fake_load_simple),
            patch.object(hx, "_build", side_effect=_inline_build),
            patch.object(hx, "_storage", return_value=self.storage),
        ]
        for p in self.patches:
            p.start()
        hx._sem = None

    def tearDown(self):
        for p in self.patches:
            p.stop()

    def _run(self, coro):
        return asyncio.run(coro)

    def test_create_runs_to_done_and_reuses(self):
        async def go():
            job, reused = await hx.create_job(self.redis, _params(), "u1")
            self.assertFalse(reused)
            await asyncio.gather(*list(hx._tasks))
            done = await hx.get_job(self.redis, job["job_id"])
            again, reused2 = await hx.create_job(self.redis, _params(), "u2")
            return done, again, reused2

        done, again, reused2 = self._run(go())
        self.assertEqual(done["status"], "done")
        self.assertEqual(done["progress"], 100)
        self.assertEqual(done["row_count"], 3)
        self.assertIn(done["object_key"], self.storage.objs)
        wb = load_workbook(io.BytesIO(self.storage.objs[done["object_key"]]))
        self.assertIn("汇总", wb.sheetnames)
        self.assertTrue(reused2)
        self.assertEqual(again["job_id"], done["job_id"])

    def test_running_job_is_reused_but_different_params_not(self):
        async def go():
            gate = asyncio.Event()

            async def slow(redis, job):
                await gate.wait()

            a, _ = await hx.create_job(self.redis, _params(), None, runner=slow)
            b, rb = await hx.create_job(self.redis, _params(), None, runner=slow)
            c, rc = await hx.create_job(self.redis, _params(sort="land_id"), None, runner=slow)
            gate.set()
            await asyncio.gather(*list(hx._tasks))
            return a, b, rb, c, rc

        a, b, rb, c, rc = self._run(go())
        self.assertTrue(rb)
        self.assertEqual(a["job_id"], b["job_id"])
        self.assertFalse(rc)
        self.assertNotEqual(a["job_id"], c["job_id"])

    def test_stale_heartbeat_reports_failed_and_is_not_reused(self):
        async def go():
            async def noop(redis, job):
                return None

            job, _ = await hx.create_job(self.redis, _params(), None, runner=noop)
            await asyncio.gather(*list(hx._tasks))
            import json

            raw = json.loads(self.redis.d[f"harvest-export:job:{job['job_id']}"])
            raw["heartbeat_at"] = time.time() - hx.STALE_SEC - 5
            self.redis.d[f"harvest-export:job:{job['job_id']}"] = json.dumps(raw)
            got = await hx.get_job(self.redis, job["job_id"])
            new, reused = await hx.create_job(self.redis, _params(), None, runner=noop)
            await asyncio.gather(*list(hx._tasks))
            return got, reused

        got, reused = self._run(go())
        self.assertEqual(got["status"], "failed")
        self.assertIn("中断", got["error"])
        self.assertFalse(reused)

    def test_failure_is_recorded(self):
        async def boom(filters):
            raise RuntimeError("db down")

        async def go():
            with patch.object(hx, "_load", side_effect=boom):
                job, _ = await hx.create_job(self.redis, _params(), None)
                await asyncio.gather(*list(hx._tasks))
            return await hx.get_job(self.redis, job["job_id"])

        job = self._run(go())
        self.assertEqual(job["status"], "failed")
        self.assertIn("db down", job["error"])

    def test_chunked_build_roundtrip(self):
        from concurrent.futures import ThreadPoolExecutor
        import os
        import tempfile

        async def go():
            parcels, daily = await _fake_load_simple(F)
            fd, path = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            with patch.object(hx, "_PICKLE_CHUNK", 2), \
                 patch.object(hx, "_get_executor", return_value=ThreadPoolExecutor(1)):
                size = await hx._build(parcels, daily, hr.summarize(parcels), F, True, path)
            with open(path, "rb") as fh:
                data = fh.read()
            os.unlink(path)
            return size, data, os.path.exists(path + ".in")

        size, data, leftover = self._run(go())
        self.assertEqual(size, len(data))
        self.assertFalse(leftover)
        ws = load_workbook(io.BytesIO(data))["汇总"]
        self.assertTrue(any("地块2" == c.value for row in ws.iter_rows() for c in row))

    def test_cleanup_old_files(self):
        now = datetime(2026, 10, 13, 10, tzinfo=timezone(timedelta(hours=8)))
        for k in ["20261010/a.xlsx", "20261012/b.xlsx", "20261013/c.xlsx"]:
            self.storage.objs[hx.KEY_PREFIX + k] = b"x"
        n = hx.cleanup_old_files(self.storage, now)
        self.assertEqual(n, 1)
        self.assertEqual(self.storage.deleted, [hx.KEY_PREFIX + "20261010/a.xlsx"])


async def _fake_load_simple(filters):
    lands = [_land(str(i)) for i in range(3)]
    return hr.build_report(filters, lands, [], [], {}, {}, date(2026, 10, 10))


class ExcludeParamsTests(unittest.TestCase):
    def test_hash_includes_exclusion_and_roundtrip(self):
        f = hr.ReportFilters(date_from=date(2026, 8, 1), date_to=date(2026, 10, 10))
        a = hx.params_of(f, "combined_pct", "desc", True)
        f2 = hr.ReportFilters(date_from=date(2026, 8, 1), date_to=date(2026, 10, 10),
                              exclude_base_ids=["4", "62"])
        b = hx.params_of(f2, "combined_pct", "desc", True)
        self.assertNotEqual(hx.params_hash(a), hx.params_hash(b))
        self.assertEqual(hx.filters_from(b).exclude_base_ids, ["4", "62"])
        self.assertEqual(hx.filters_from({"from": "2026-08-01", "to": "2026-10-10"}).exclude_base_ids, [])


class ExportJobRouterTests(unittest.TestCase):
    def test_endpoints(self):
        from app.middleware.auth import require_roles  # noqa: F401
        from app.routers import harvest_report as router_mod

        app = FastAPI()
        app.include_router(router_mod.router, prefix="/v1")
        app.state.redis_client = FakeRedis()
        storage = FakeStorage()
        for dep in list(app.dependency_overrides):
            pass
        app.dependency_overrides[router_mod._reader] = lambda: type("C", (), {"user": None})()

        async def fake_create(redis, params, created_by, runner=None):
            job = {"job_id": "j1", "status": "done", "progress": 100, "stage": "完成", "error": None,
                   "params": params, "file_name": "harvest-report_x.xlsx", "object_key": "k",
                   "size_bytes": 3, "row_count": 1, "created_by": None, "created_at": time.time(),
                   "started_at": None, "finished_at": time.time(), "heartbeat_at": time.time()}
            await hx._save(redis, job)
            return job, False

        storage.objs["k"] = b"abc"
        with patch.object(hx, "create_job", side_effect=fake_create), \
             patch.object(hx, "_storage", return_value=storage):
            c = TestClient(app)
            r = c.post("/v1/agri/harvest-report/export-jobs?from=2026-08-01&to=2026-10-10&status=lt30")
            self.assertEqual(r.status_code, 202, r.text)
            self.assertEqual(r.json()["params"]["status"], ["lt30"])
            self.assertEqual(c.get("/v1/agri/harvest-report/export-jobs/j1").json()["status"], "done")
            d = c.get("/v1/agri/harvest-report/export-jobs/j1/download")
            self.assertEqual(d.status_code, 200)
            self.assertEqual(d.content, b"abc")
            self.assertIn("attachment", d.headers["content-disposition"])
            self.assertEqual(c.get("/v1/agri/harvest-report/export-jobs/nope").status_code, 404)
            self.assertIn("62", r.json()["params"]["exclude_base_ids"])
            bad = c.post("/v1/agri/harvest-report/export-jobs?status=xx")
            self.assertEqual(bad.status_code, 422)


if __name__ == "__main__":
    unittest.main()
