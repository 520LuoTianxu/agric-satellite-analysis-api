"""全部地块收获进度报表：组装逻辑、筛选排序、Excel 结构与接口参数。"""

from __future__ import annotations

import asyncio
import io
import unittest
from datetime import date, datetime
from unittest.mock import AsyncMock, MagicMock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services import harvest_report as hr
from app.services.harvest_report_xlsx import CN_TZ, build_workbook

F = hr.ReportFilters(date_from=date(2026, 8, 1), date_to=date(2026, 10, 10))
TODAY = date(2026, 10, 10)


def _land(lid: str, area: float = 100.0, **kw) -> dict:
    return {
        "land_id": lid,
        "land_name": f"地块{lid}",
        "group_id": "12268",
        "group_name": "测试组",
        "crop_type": "corn",
        "area_mu": area,
        **kw,
    }


def _row(
    lid: str,
    d: str,
    combined: float,
    harvested: float | None = None,
    season: str = "2026-05-01",
    **kw,
) -> dict:
    h = combined if harvested is None else harvested
    return {
        "land_id": lid,
        "obs_date": date.fromisoformat(d),
        "status": "harvesting",
        "harvested_pct": h,
        "suspected_harvest_pct": round(combined - h, 1),
        "harvested_or_suspected_pct": combined,
        "newly_harvested_pct": 1.0,
        "harvested_area_mu": None,
        "parcel_area_mu": 100.0,
        "valid_pct": 95.0,
        "season_start": date.fromisoformat(season),
        "confidence": 0.8,
        "confidence_level": "high",
        "confirmed": True,
        "method_version": "s2s1_residue_monotonic_v4",
        **kw,
    }


def _cov(f: str, t: str) -> dict:
    return {
        "computed_from": date.fromisoformat(f),
        "computed_to": date.fromisoformat(t),
        "min_save_pct": 3,
    }


def _build(lands, rows=(), prev=(), cov=None, latest=None, filters=F):
    return hr.build_report(filters, lands, rows, prev, cov or {}, latest or {}, TODAY)


class BuildReportTests(unittest.TestCase):
    def test_latest_state_first_date_and_increment(self):
        parcels, daily = _build(
            [_land("1")],
            [
                _row("1", "2026-09-10", 10, 5),
                _row("1", "2026-09-20", 40, 30),
                _row("1", "2026-10-01", 95, 90),
            ],
            latest={"1": date(2026, 10, 6)},
        )
        p = parcels[0]
        self.assertEqual(p["first_harvest_date"], date(2026, 9, 10))
        self.assertEqual(p["latest_date"], date(2026, 10, 1))
        self.assertEqual(
            (p["harvested_pct"], p["suspected_pct"], p["combined_pct"]),
            (90.0, 5.0, 95.0),
        )
        self.assertEqual(p["newly_combined_pct"], 55.0)
        self.assertEqual(p["harvested_area_mu"], 90.0)  # 未存面积时按 面积×已收获%
        self.assertEqual(p["combined_area_mu"], 95.0)
        self.assertEqual(p["days_since_last_image"], 4)
        self.assertEqual(p["bucket"], hr.BUCKET_GE90)
        self.assertEqual(p["coverage"], "stored")
        # 首期无上一条：新增即本期合计（之前均 ≤ 门槛）
        self.assertEqual([r["newly_combined_pct"] for r in daily], [10.0, 30.0, 55.0])

    def test_prev_row_before_range_used_for_increment(self):
        _p, daily = _build(
            [_land("1")],
            [_row("1", "2026-08-05", 50)],
            prev=[_row("1", "2026-07-30", 20)],
        )
        self.assertEqual(daily[0]["newly_combined_pct"], 30.0)

    def test_new_season_resets_first_date_and_increment(self):
        parcels, daily = _build(
            [_land("1")],
            [
                _row("1", "2026-08-05", 100, season="2025-05-01"),
                _row("1", "2026-09-25", 20, season="2026-05-01"),
            ],
        )
        self.assertEqual(parcels[0]["first_harvest_date"], date(2026, 9, 25))
        self.assertEqual(daily[1]["newly_combined_pct"], 20.0)

    def test_coverage_without_rows_is_computed_zero(self):
        parcels, _ = _build([_land("1")], cov={"1": _cov("2026-01-01", "2026-10-10")})
        p = parcels[0]
        self.assertTrue(p["computed"])
        self.assertEqual(p["combined_pct"], 0.0)
        self.assertEqual(p["bucket"], hr.BUCKET_NONE)
        self.assertEqual(p["status"], "growing")

    def test_coverage_stale_but_no_newer_image_is_full(self):
        parcels, _ = _build(
            [_land("1")],
            cov={"1": _cov("2026-01-01", "2026-10-01")},
            latest={"1": date(2026, 9, 30)},
        )
        self.assertEqual(parcels[0]["bucket"], hr.BUCKET_NONE)

    def test_partial_or_missing_coverage_is_not_computed(self):
        parcels, _ = _build(
            [_land("1"), _land("2")],
            cov={"1": _cov("2026-09-01", "2026-10-10")},
        )
        self.assertEqual([p["coverage"] for p in parcels], ["partial", "none"])
        for p in parcels:
            self.assertFalse(p["computed"])
            self.assertIsNone(p["combined_pct"])
            self.assertEqual(p["bucket"], hr.BUCKET_NOT_COMPUTED)

    def test_rows_without_coverage_marker_count_as_computed(self):
        parcels, _ = _build([_land("1")], [_row("1", "2026-09-01", 25)])
        self.assertTrue(parcels[0]["computed"])
        self.assertEqual(parcels[0]["bucket"], hr.BUCKET_LT30)

    def test_legacy_row_without_combined_uses_harvested(self):
        r = _row("1", "2026-09-01", 40)
        r["harvested_or_suspected_pct"] = None
        r["suspected_harvest_pct"] = None
        parcels, _ = _build([_land("1")], [r])
        self.assertEqual(
            (parcels[0]["combined_pct"], parcels[0]["suspected_pct"]), (40.0, 0.0)
        )


class FilterSortSummaryTests(unittest.TestCase):
    def setUp(self):
        self.parcels, _ = _build(
            [_land("1", 100), _land("2", 200), _land("3", 50), _land("4", 10)],
            [_row("1", "2026-09-01", 95), _row("2", "2026-09-01", 40)],
            cov={"3": _cov("2026-01-01", "2026-10-10")},
        )

    def test_buckets_and_min_pct(self):
        f = hr.ReportFilters(F.date_from, F.date_to, buckets=["ge90", "not_computed"])
        self.assertEqual(
            {p["land_id"] for p in hr.apply_filters(self.parcels, f)}, {"1", "4"}
        )
        f = hr.ReportFilters(F.date_from, F.date_to, min_pct=40)
        self.assertEqual(
            {p["land_id"] for p in hr.apply_filters(self.parcels, f)}, {"1", "2"}
        )

    def test_sort_nulls_last_both_orders(self):
        for order in ("asc", "desc"):
            ids = [
                p["land_id"]
                for p in hr.sort_parcels(self.parcels, "combined_pct", order)
            ]
            self.assertEqual(ids[-1], "4")
        self.assertEqual(
            [
                p["land_id"]
                for p in hr.sort_parcels(self.parcels, "combined_pct", "desc")
            ],
            ["1", "2", "3", "4"],
        )
        # 未知字段回退到 combined_pct
        self.assertEqual(
            hr.sort_parcels(self.parcels, "drop table", "asc")[0]["land_id"], "3"
        )

    def test_summary(self):
        s = hr.summarize(self.parcels)
        self.assertEqual(s["parcel_count"], 4)
        self.assertEqual(s["computed_count"], 3)
        self.assertEqual(s["total_area_mu"], 360.0)
        self.assertEqual(s["combined_area_mu"], 175.0)  # 95 + 80 + 0
        self.assertEqual(s["avg_combined_pct"], 45.0)  # (95+40+0)/3，未计算不计入
        self.assertEqual(s["area_weighted_combined_pct"], 50.0)  # 175/350
        self.assertEqual(
            s["bucket_counts"],
            {"none": 1, "lt30": 0, "30_90": 1, "ge90": 1, "not_computed": 1},
        )


class WorkbookTests(unittest.TestCase):
    def _wb(self, carry=True):
        parcels, daily = _build(
            [_land("1"), _land("2"), _land("3")],
            [
                _row("1", "2026-09-01", 10),
                _row("1", "2026-09-10", 60),
                _row("2", "2026-09-05", 30),
            ],
            cov={"3": _cov("2026-01-01", "2026-10-10")},
        )
        data = build_workbook(
            parcels,
            daily,
            hr.summarize(parcels),
            F,
            datetime(2026, 10, 10, 16, 30, tzinfo=CN_TZ),
            carry,
        )
        return load_workbook(io.BytesIO(data))

    def test_sheets_headers_freeze_filter(self):
        wb = self._wb()
        self.assertEqual(wb.sheetnames, ["汇总", "逐日明细", "透视", "说明"])
        s = wb["汇总"]
        self.assertEqual(s["A7"].value, "地块ID")
        self.assertEqual(s["A5"].value, 3)  # 地块数
        self.assertEqual(s.freeze_panes, "C8")
        self.assertTrue(s.auto_filter.ref.startswith("A7:"))
        self.assertEqual(s.max_row, 10)
        self.assertGreater(len(s.conditional_formatting), 0)
        d = wb["逐日明细"]
        self.assertEqual(d.max_row, 4)
        self.assertEqual(d.freeze_panes, "C2")
        self.assertIn("2026-10-10 16:30", wb["说明"]["B1"].value)

    def test_pivot_carry_forward(self):
        p = self._wb()["透视"]
        self.assertEqual(
            [c.value.date() for c in p[1][3:]],
            [date(2026, 9, 1), date(2026, 9, 5), date(2026, 9, 10)],
        )
        rows = {r[0].value: [c.value for c in r[3:]] for r in p.iter_rows(min_row=2)}
        self.assertEqual(rows["1"], [10, 10, 60])  # 09-05 沿用
        self.assertEqual(rows["2"], [None, 30, 30])
        self.assertEqual(rows["3"], [None, None, None])
        self.assertTrue(p.cell(row=2, column=5).font.i)
        plain = {
            r[0].value: [c.value for c in r[3:]]
            for r in self._wb(carry=False)["透视"].iter_rows(min_row=2)
        }
        self.assertEqual(plain["1"], [10, None, 60])


class EndpointTests(unittest.TestCase):
    def _client(self):
        from app.core.database import get_db
        from app.routers import harvest_report as router_mod

        app = FastAPI()
        app.include_router(router_mod.router, prefix="/v1")
        app.dependency_overrides[get_db] = lambda: MagicMock()
        app.dependency_overrides[router_mod._reader] = lambda: MagicMock()
        return app, router_mod

    def _data(self):
        return _build(
            [_land("1"), _land("2"), _land("3")],
            [_row("1", "2026-09-01", 95), _row("2", "2026-09-01", 20)],
        )

    def test_json_paginates_filters_and_passes_params(self):
        app, mod = self._client()
        load = AsyncMock(return_value=self._data())
        with patch.object(mod.hr, "load_report", load):
            r = TestClient(app).get(
                "/v1/agri/harvest-report",
                params={
                    "from": "2026-08-01",
                    "to": "2026-10-10",
                    "group_id": "12268",
                    "crop": "corn",
                    "status": "ge90,lt30",
                    "page_size": 1,
                },
            )
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertEqual(body["total"], 2)
        self.assertEqual([i["land_id"] for i in body["items"]], ["1"])
        self.assertEqual(body["items"][0]["latest_date"], "2026-09-01")
        self.assertEqual(body["summary"]["parcel_count"], 2)
        # 下拉选项基于分档/门槛筛选前的地块
        self.assertEqual(
            body["facets"],
            {
                "groups": [{"id": "12268", "name": "测试组", "count": 3}],
                "crops": [{"id": "corn", "count": 3}],
            },
        )
        f = load.await_args.args[1]
        self.assertEqual(
            (f.group_id, f.crop, f.buckets), ("12268", "corn", ["ge90", "lt30"])
        )

    def test_validation(self):
        app, mod = self._client()
        c = TestClient(app)
        with patch.object(mod.hr, "load_report", AsyncMock(return_value=([], []))):
            self.assertEqual(
                c.get("/v1/agri/harvest-report?status=bogus").status_code, 422
            )
            self.assertEqual(
                c.get(
                    "/v1/agri/harvest-report?from=2026-10-02&to=2026-10-01"
                ).status_code,
                422,
            )
            self.assertEqual(
                c.get(
                    "/v1/agri/harvest-report?from=2024-01-01&to=2026-10-01"
                ).status_code,
                422,
            )
            self.assertEqual(c.get("/v1/agri/harvest-report").status_code, 200)

    def test_export_xlsx(self):
        app, mod = self._client()
        with patch.object(mod.hr, "load_report", AsyncMock(return_value=self._data())):
            r = TestClient(app).get(
                "/v1/agri/harvest-report/export.xlsx?from=2026-08-01&to=2026-10-10&min_pct=50"
            )
        self.assertEqual(r.status_code, 200)
        self.assertIn(
            "harvest-report_20260801-20261010.xlsx", r.headers["content-disposition"]
        )
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb["汇总"]["A5"].value, 1)
        self.assertEqual(wb["逐日明细"].max_row, 2)  # 只含筛选后地块


class LoadReportTests(unittest.TestCase):
    def test_reads_all_queries_and_tolerates_missing_coverage_table(self):
        class _Res:
            def __init__(self, rows):
                self._rows = rows

            def fetchall(self):
                return self._rows

        def _r(**kw):
            m = MagicMock(**kw)
            m._mapping = kw
            return m

        db = MagicMock()

        async def execute(stmt, params):
            sql = str(stmt)
            if "FROM agric_satellite.land_parcels" in sql:
                return _Res([_r(**_land("1"))])
            if "parcel_harvest_progress_coverage" in sql:
                raise RuntimeError("relation does not exist")
            if "DISTINCT ON" in sql:
                return _Res([])
            if "parcel_harvest_progress" in sql:
                return _Res([_r(**_row("1", "2026-09-01", 50))])
            if "parcel_scene_products" in sql:
                return _Res([_r(land_id="1", obs_date=date(2026, 10, 6))])
            raise AssertionError(sql)

        db.execute = execute
        nested = MagicMock()
        nested.__aenter__ = AsyncMock()
        nested.__aexit__ = AsyncMock(return_value=False)
        db.begin_nested = MagicMock(return_value=nested)
        parcels, daily = asyncio.run(hr.load_report(db, F))
        self.assertEqual(parcels[0]["combined_pct"], 50.0)
        self.assertEqual(parcels[0]["latest_obs_date"], date(2026, 10, 6))
        self.assertEqual(len(daily), 1)


if __name__ == "__main__":
    unittest.main()
